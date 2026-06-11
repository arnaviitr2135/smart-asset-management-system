from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("CultTrack_AI_Project_Documentation_Updated.docx")


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "6B7280"
LIGHT = "F2F4F7"
SOFT_BLUE = "E8EEF5"
GREEN = "DFF3E8"
GOLD = "FFF3CD"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=INK):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(hdr[i], header, bold=True, color="111827")
        set_cell_shading(hdr[i], LIGHT)
    for row_data in rows:
        row = table.add_row().cells
        for i, value in enumerate(row_data):
            set_cell_text(row[i], str(value))
    if widths:
        set_table_widths(table, widths)
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_callout(doc, title, body, fill=SOFT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.3)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    r.font.size = Pt(10)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor.from_string(INK)
    doc.add_paragraph()


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)


def add_header_footer(doc):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("CultTrack AI Project Documentation")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Smart Asset Management and Resource Allocation Platform")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)
    add_header_footer(doc)

    # Cover page
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("CultTrack AI")
    r.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Smart Asset Management and Resource Allocation Platform")
    r.font.size = Pt(17)
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("Updated Project Documentation")
    r.bold = True
    r.font.size = Pt(14)

    add_callout(
        doc,
        "Project Scope",
        "A full-stack web application for the Cultural Council of IIT Roorkee that centralizes inventory, booking requests, approvals, handovers, returns, audit logs, email notifications, password recovery, QR-assisted operations, analytics, and demand forecasting.",
        fill=SOFT_BLUE,
    )

    meta_rows = [
        ("Submitted By", "Yash Sharma (24116108), Arnav Saini (24116015)"),
        ("Department", "Electronics and Communication Engineering, IIT Roorkee"),
        ("Repository", "https://github.com/arnaviitr2135/smart-asset-management-system"),
        ("Frontend", "https://frontend-seven-lyart-83.vercel.app"),
        ("Backend API", "https://smart-asset-management-system-oiyh.onrender.com"),
        ("Local Run", "docker compose up --build"),
        ("Last Updated", "12 June 2026"),
    ]
    add_table(doc, ["Field", "Details"], meta_rows, widths=[1.8, 4.5])
    doc.add_page_break()

    doc.add_heading("1. Problem Understanding", level=1)
    doc.add_heading("1.1 Background and Context", level=2)
    doc.add_paragraph(
        "The Cultural Council of IIT Roorkee manages a large pool of high-value shared assets used by sections, societies, and event teams. These assets include cameras, studio lights, recording equipment, sound systems, costumes, props, and event infrastructure. Because the same equipment moves between many students and events, the Council needs reliable visibility into who has an item, when it is due back, whether it is damaged, and whether it is available for a future event."
    )
    doc.add_heading("1.2 Current Manual Workflow", level=2)
    add_bullets(
        doc,
        [
            "Separate spreadsheets maintained by different coordinators.",
            "Physical logbooks for check-ins and check-outs.",
            "Informal handovers coordinated through messages, calls, and word of mouth.",
            "No reliable real-time view of availability across upcoming dates.",
            "No permanent digital trail for accountability, condition, and return history.",
        ],
    )
    doc.add_heading("1.3 Operational Pain Points", level=2)
    add_bullets(
        doc,
        [
            "Overbooking during major events because overlapping requests are hard to detect manually.",
            "Stock visibility gaps where an item appears available in one sheet but is physically checked out.",
            "Lost accountability for damaged or missing accessories because handovers are not consistently recorded.",
            "Slow issue and return operations when admins must search and type asset names manually.",
            "No analytics for utilization, shortage prediction, or purchase planning.",
        ],
    )

    doc.add_heading("2. Proposed Solution", level=1)
    doc.add_paragraph(
        "CultTrack AI replaces the fragmented manual system with a centralized full-stack portal. It gives members a clean way to browse inventory and request assets, while admins get a control desk for approvals, issue, return, inventory management, health reports, notifications, analytics, and audit history."
    )
    add_table(
        doc,
        ["Objective", "How CultTrack AI Solves It"],
        [
            ("Prevent overbooking", "Date-aware availability checks compare requested quantity against approved reservations and active issued allocations without double-counting."),
            ("Improve accountability", "Every booking, approval, issue, return, health report, and important action is persisted in the database and audit log."),
            ("Speed up operations", "Admin desk, QR scan simulation, and focused check-in views reduce manual searching during handover and return."),
            ("Improve communication", "In-app notifications and email notifications inform users and admins about requests, approvals, returns, and password resets."),
            ("Support planning", "Analytics and demand forecasting expose utilization and predicted shortages."),
        ],
        widths=[1.8, 4.5],
    )

    doc.add_heading("3. Complete Feature List", level=1)
    add_table(
        doc,
        ["Area", "Features", "Primary Users"],
        [
            ("Authentication", "JWT login, registration, role-based sessions, profile restore, protected API routes.", "All users"),
            ("Password Recovery", "Forgot-password request, secure reset token, reset-password screen, generic responses for account privacy.", "All users"),
            ("Email", "Resend API support, Gmail SMTP fallback, password reset emails, booking and return notifications.", "All users"),
            ("Inventory Catalog", "Search, category filters, live shelf counts, active/maintenance/damaged state, responsive asset cards.", "Members"),
            ("Date Availability", "Dedicated availability endpoint and modal count for selected dates; prevents mismatch between shelf count and request availability.", "Members"),
            ("Booking Requests", "Quantity, date range, purpose, validation, pending request submission, admin review queue.", "Members, admins"),
            ("Admin Approval", "Approve/reject requests, reserve physical stock on approval, notify users and admins.", "Admins"),
            ("Issue/Handover", "Convert approved bookings into allocations, record issuing admin, due date, and active loan status.", "Admins"),
            ("Return Requests", "Member return button, admin-initiated return request, return request log, user confirmation, persisted statuses.", "Members, admins"),
            ("Final Check-in", "Admin check-in with condition on return, return record, inventory restoration, maintenance updates for damaged items.", "Admins"),
            ("Health Logs", "Manual health reports, damaged condition tracking, maintenance state handling.", "Admins"),
            ("QR Simulation", "QR code payloads and scan simulation for fast asset lookup and operational workflows.", "Admins"),
            ("Notifications", "Unread count, notification center, mark-read action, booking and return status alerts.", "All users"),
            ("Audit Trail", "Permanent action history for approvals, rejections, issue, return, inventory changes, health reports, and password resets.", "Admins"),
            ("Analytics", "Dashboard metrics, utilization rates, active loans, overdue counts, and forecast data.", "Admins"),
            ("Demand Forecasting", "Weekly weighted moving-average style forecast and shortage risk detection.", "Admins"),
            ("Deployment", "Docker local stack, Vercel frontend, Render backend, Neon PostgreSQL production database.", "Developers"),
        ],
        widths=[1.4, 3.9, 1.0],
    )

    doc.add_heading("4. User Roles and Workflows", level=1)
    doc.add_heading("4.1 Society Member Workflow", level=2)
    add_numbered(
        doc,
        [
            "Register or log in using a member account.",
            "Browse the inventory catalog by search and category.",
            "Open an asset request modal and select quantity, start date, end date, and purpose.",
            "Review the date-specific availability count shown in the modal.",
            "Submit the booking request and wait for admin approval.",
            "Track pending, approved, issued, rejected, and returned items in Bookings and Loans.",
            "Use Return Asset when ready to return an issued item.",
            "Confirm admin return requests from the Return Request Log.",
            "Receive in-app and email notifications for booking and return status changes.",
        ],
    )
    doc.add_heading("4.2 Council Admin Workflow", level=2)
    add_numbered(
        doc,
        [
            "Log in with an admin account.",
            "Create, edit, or remove inventory assets.",
            "Review pending booking requests in the Council Control Desk.",
            "Approve or reject requests after checking stock and event context.",
            "Issue approved assets during handover and create allocation records.",
            "Monitor active outstanding loans, due dates, and overdue items.",
            "Request users to return assets when needed.",
            "Complete check-in by recording returned condition and notes.",
            "Log damaged equipment and move it to maintenance if required.",
            "Review audit logs, health logs, analytics, and demand forecasts.",
        ],
    )

    doc.add_heading("5. System Architecture", level=1)
    add_table(
        doc,
        ["Layer", "Technology", "Responsibility"],
        [
            ("Frontend", "React, Vite, TypeScript, Tailwind CSS, Lucide icons, Recharts", "User interface for inventory, booking, loans, admin desk, analytics, and authentication."),
            ("Backend", "Node.js, Express, TypeScript", "REST API, validation, role authorization, booking logic, notifications, and business workflows."),
            ("ORM", "Prisma", "Type-safe database access and schema management."),
            ("Database", "PostgreSQL", "Relational storage for users, assets, bookings, allocations, returns, health reports, notifications, audit logs, and reset tokens."),
            ("Local Runtime", "Docker Compose", "Runs frontend, backend, and PostgreSQL locally with seeded demo data."),
            ("Production", "Vercel, Render, Neon", "Frontend hosting, backend Docker service, and managed PostgreSQL database."),
        ],
        widths=[1.2, 2.1, 3.0],
    )
    add_callout(
        doc,
        "Deployment Summary",
        "The frontend is deployed on Vercel, the backend API is deployed on Render, and production data is stored in Neon PostgreSQL. Local development uses Docker Compose and Postgres 16.",
        fill=GREEN,
    )

    doc.add_heading("6. Database Schema", level=1)
    add_table(
        doc,
        ["Model", "Purpose", "Important Fields"],
        [
            ("User", "Stores all member and admin accounts.", "email, passwordHash, fullName, role"),
            ("Asset", "Represents inventory items.", "name, category, quantityAvailable, totalQuantity, status, qrCodeUrl"),
            ("Booking", "Initial member request for an asset.", "userId, assetId, quantity, startDate, endDate, purpose, status"),
            ("Allocation", "Active issued loan after admin handover.", "bookingId, assetId, userId, issuedById, dueDate, status"),
            ("Return", "Completed return/check-in record.", "allocationId, receivedById, returnedAt, conditionOnReturn, notes"),
            ("ReturnRequest", "Persisted handover/return coordination request.", "allocationId, userId, requestedById, source, status, notes"),
            ("AssetHealthReport", "Maintenance and damage reporting.", "assetId, reportedById, condition, notes"),
            ("Notification", "In-app notification feed.", "userId, title, message, type, isRead"),
            ("AuditLog", "Administrative and security action history.", "userId, action, details, createdAt"),
            ("PasswordResetToken", "Secure password reset tokens.", "userId, tokenHash, expiresAt"),
        ],
        widths=[1.3, 2.2, 2.8],
    )
    doc.add_paragraph(
        "Core relationship flow: a User creates a Booking. After admin approval and handover, the Booking becomes an Allocation. When the item is returned, a Return record is created and the Allocation status changes to RETURNED. ReturnRequest records coordinate user/admin return handover before final check-in."
    )

    doc.add_heading("7. API Overview", level=1)
    add_table(
        doc,
        ["Route Group", "Key Endpoints", "Purpose"],
        [
            ("Auth", "POST /api/v1/auth/register, /login, /forgot-password, /reset-password, GET /me", "Account creation, login, password recovery, and session validation."),
            ("Assets", "GET /api/v1/assets, GET /:id, GET /:id/availability, POST/PUT/DELETE /:id", "Inventory listing, date-specific availability, and admin asset management."),
            ("Bookings", "POST /api/v1/bookings, GET /api/v1/bookings, PUT /:id/approve, PUT /:id/reject", "Request submission and admin review."),
            ("Allocations", "GET /api/v1/allocations, POST /issue, POST /return", "Issue, active loans, and return check-in."),
            ("Return Requests", "POST /request-return, GET /return-requests, POST /return-requests/:id/respond", "Member/admin return coordination."),
            ("Health", "POST /api/v1/allocations/health, GET /api/v1/allocations/health", "Damage and maintenance reporting."),
            ("Analytics", "GET /api/v1/analytics, GET /api/v1/analytics/forecast", "Metrics and demand forecasting."),
            ("Audit", "GET /api/v1/audit", "Admin audit trail."),
            ("Notifications", "GET /api/v1/notifications, PUT /:id/read", "In-app notification center."),
            ("Health Check", "GET /healthz", "Backend and database connectivity check."),
        ],
        widths=[1.35, 3.35, 1.6],
    )
    doc.add_paragraph("Most API routes require an Authorization: Bearer <jwt_token> header. Admin-only routes additionally enforce the ADMIN role.")

    doc.add_heading("8. Booking and Availability Logic", level=1)
    doc.add_paragraph(
        "Booking correctness is one of the most important parts of the system. The backend validates quantity, date format, non-past start dates, and overlapping approved reservations. It also checks current physical shelf stock before accepting a request."
    )
    add_bullets(
        doc,
        [
            "The catalog card shows current shelf stock, such as 4 / 8 in stock.",
            "The booking modal calls /api/v1/assets/:id/availability for the selected start and end dates.",
            "The modal shows Available for selected dates, preventing confusion between shelf stock and date-specific availability.",
            "Approved bookings reserve stock. Issued allocations are linked to those bookings, so the backend avoids double-counting the same loan.",
            "Returns restore quantityAvailable and mark the allocation as RETURNED.",
        ],
    )

    doc.add_heading("9. Notifications, Email, and Password Reset", level=1)
    add_table(
        doc,
        ["Capability", "Behavior"],
        [
            ("In-app notifications", "Users see unread counts and status messages for bookings, approvals, issue, returns, and admin requests."),
            ("Email provider", "Resend API is the recommended provider; Gmail SMTP is available as fallback."),
            ("Password reset", "A reset token is created, emailed to the user, and expires after a short validity window."),
            ("Privacy behavior", "Forgot-password responses do not reveal whether an email is registered."),
            ("Operational logs", "Render logs indicate whether email is skipped, sent through Resend, or connected through SMTP."),
        ],
        widths=[1.8, 4.5],
    )

    doc.add_heading("10. Analytics and Forecasting", level=1)
    doc.add_paragraph(
        "The admin analytics dashboard turns operational data into planning insight. It summarizes inventory, booking, allocation, and utilization data, then uses a lightweight weekly weighted moving-average style forecast to identify demand risk without requiring a heavy machine learning deployment."
    )
    add_bullets(
        doc,
        [
            "Dashboard metrics for assets, bookings, issued loans, overdue loans, and utilization.",
            "Forecast endpoint that reads past booking and allocation activity.",
            "Shortage risk flagging when projected demand approaches available inventory.",
            "Recharts-based frontend visualizations for admin review.",
        ],
    )

    doc.add_heading("11. Local Setup and Deployment", level=1)
    doc.add_heading("11.1 Docker Setup", level=2)
    add_callout(doc, "Command", "git clone https://github.com/arnaviitr2135/smart-asset-management-system.git\ncd smart-asset-management-system\ndocker compose up --build", fill=LIGHT)
    add_table(
        doc,
        ["Service", "Local URL", "Notes"],
        [
            ("Frontend", "http://localhost:5173", "Vite React application."),
            ("Backend", "http://localhost:5000", "Express API."),
            ("Health", "http://localhost:5000/healthz", "Checks API and database connectivity."),
            ("Database", "localhost:5432", "PostgreSQL container."),
        ],
        widths=[1.4, 2.2, 2.7],
    )
    doc.add_heading("11.2 Production Setup", level=2)
    add_table(
        doc,
        ["Component", "Provider", "Production Value"],
        [
            ("Frontend", "Vercel", "https://frontend-seven-lyart-83.vercel.app"),
            ("Backend", "Render", "https://smart-asset-management-system-oiyh.onrender.com"),
            ("Database", "Neon", "PostgreSQL connection string stored in Render env vars."),
            ("Email", "Resend / SMTP", "RESEND_API_KEY or SMTP_USER and SMTP_PASS in Render."),
        ],
        widths=[1.4, 1.5, 3.4],
    )

    doc.add_heading("12. Environment Variables", level=1)
    add_table(
        doc,
        ["Variable", "Used By", "Purpose"],
        [
            ("DATABASE_URL", "Backend", "PostgreSQL/Neon connection string."),
            ("JWT_SECRET", "Backend", "Signs and verifies JWT tokens."),
            ("PORT", "Backend", "API port, usually 5000 locally and Render-assigned in production."),
            ("FRONTEND_URL", "Backend", "Builds reset links and controls allowed frontend origin."),
            ("RESEND_API_KEY", "Backend", "Recommended email provider API key."),
            ("RESEND_FROM", "Backend", "Sender identity for Resend emails."),
            ("SMTP_USER / SMTP_PASS", "Backend", "Gmail SMTP fallback credentials."),
            ("VITE_API_URL", "Frontend", "Base URL for API calls from the Vercel frontend."),
        ],
        widths=[1.7, 1.2, 3.4],
    )

    doc.add_heading("13. Security and Reliability", level=1)
    add_bullets(
        doc,
        [
            "Passwords are stored as bcrypt hashes, not plaintext.",
            "Protected API routes require JWT authentication.",
            "Admin-only actions are guarded by role middleware.",
            "Password reset tokens are hashed before storage and expire after a short window.",
            "Secrets must remain in environment variables and should never be committed.",
            "Audit logs preserve important operational actions for accountability.",
            "Docker and Prisma schema sync make local setup reproducible.",
        ],
    )

    doc.add_heading("14. Troubleshooting Guide", level=1)
    add_table(
        doc,
        ["Issue", "Likely Cause", "Resolution"],
        [
            ("Frontend cannot reach backend", "VITE_API_URL is wrong or Render service is asleep/down.", "Check /healthz and confirm Vercel environment variable."),
            ("Forgot password email missing", "Email not registered, Resend/SMTP misconfigured, or Render not redeployed after env change.", "Check Render logs for Password Reset, Resend, SMTP, or Email Error lines."),
            ("Wrong availability count", "Backend not deployed to latest availability fix or date-specific count differs from shelf stock.", "Deploy latest Render commit and use modal Available for selected dates count."),
            ("Render Dockerfile error", "Incorrect Dockerfile path or build context.", "Use backend/Dockerfile.render as Dockerfile path and backend as build context."),
            ("Local data looks stale", "Old Docker Postgres volume still mounted.", "Run docker compose down -v, then docker compose up --build."),
        ],
        widths=[1.7, 2.3, 2.3],
    )

    doc.add_heading("15. Future Enhancements", level=1)
    add_bullets(
        doc,
        [
            "Real QR scanner integration using device camera APIs.",
            "Calendar view for asset reservations by date.",
            "Admin export reports for semester audits and budget planning.",
            "Verified production email domain for higher deliverability.",
            "Fine-grained society/team ownership and multi-admin approval policies.",
            "Automated overdue reminders and escalation workflows.",
        ],
    )

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("End of Updated Project Documentation")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build_doc()
    print(path)
